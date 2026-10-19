"""MxSheet — port of MxLoader's MxSheet.cls worksheet layout.

    row 1: A=object structure  B=main object  C=action  D=where  E=order by  F=OSLC directives
    row 2: attribute header, one column per attribute
    row 3+: data

Header conventions (identical to MxLoader):

- `ATTR`            attribute of the main object
- `REL.ATTR`        attribute of child object REL (the OSLC child name, e.g.
                    `POLINE.ORDERQTY`); `REL.SUB.ATTR` goes a level deeper
- `[ATTR]`          an ID to *return*: never sent, filled from the response
                    (e.g. `[ASSETUID]`) — also sent in the `properties` header
- `[RESULT]`        per-row outcome ("OK" or the Maximo error message)

Any attribute name is accepted — standard or custom (`EX_CRITICALITY`,
`POLINE.EX_CONTRACTREF`): the header *is* the attribute list, which is why
MxLoader-style loads need no configuration to handle custom attributes.

Backed by either a CSV file (same row 1 / row 2 / data layout, so it can be
edited in any spreadsheet tool and diffed in git) or an .xlsx/.xlsm workbook
(read and written in place via openpyxl — install the `excel` extra).
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

HEADER_ROW = 2
START_ROW = 3
CONFIG_CELLS = ("object_service", "object_name", "action", "where", "order_by", "oslc_directives")

RESULT_HEADER = "[RESULT]"
ERROR_HEADER = "[ERROR]"

LOAD_ACTIONS = ("AddChange", "Add", "Change", "Delete", "Sync")
ACTIONS = ("Query",) + LOAD_ACTIONS
_ACTION_LOOKUP = {a.lower(): a for a in ACTIONS}


class SheetFormatError(ValueError):
    """The sheet doesn't follow the MxLoader row 1 / row 2 layout."""


@dataclass(frozen=True)
class Column:
    index: int  # 0-based column index
    header: str  # normalized (uppercase, no newlines)
    kind: str  # "attr" | "id" | "result" | "error"
    path: tuple[str, ...] = ()  # child relation chain, lowercase; () = main object
    attr: str = ""  # lowercase attribute name

    @property
    def is_data(self) -> bool:
        return self.kind == "attr"


def parse_header(index: int, raw: Any) -> Column:
    header = str(raw or "").replace("\n", "").replace("\r", "").strip().upper()
    if not header:
        raise SheetFormatError(f"empty attribute header in column {index + 1}")
    if header == RESULT_HEADER:
        return Column(index, header, "result")
    if header == ERROR_HEADER:
        return Column(index, header, "error")
    kind = "attr"
    name = header
    if header.startswith("[") and header.endswith("]"):
        kind, name = "id", header[1:-1]
    parts = [p.strip().lower() for p in name.split(".")]
    if any(not p for p in parts):
        raise SheetFormatError(f"malformed attribute header {header!r} in column {index + 1}")
    return Column(index, header, kind, tuple(parts[:-1]), parts[-1])


def normalize_action(action: str) -> str:
    try:
        return _ACTION_LOOKUP[action.strip().lower()]
    except KeyError:
        raise SheetFormatError(f"unsupported action {action!r}; expected one of {', '.join(ACTIONS)}") from None


@dataclass
class MxSheet:
    object_service: str
    object_name: str
    action: str
    columns: list[Column]
    data: list[list[Any]] = field(default_factory=list)
    where: str = ""
    order_by: str = ""
    oslc_directives: str = ""
    path: Path | None = None
    sheet_name: str | None = None

    def __post_init__(self) -> None:
        self.validate()

    # -- validation (MxSheet.Init) -------------------------------------------------
    def validate(self) -> None:
        if not self.object_service:
            raise SheetFormatError("cell A1 must name the object structure (e.g. MXASSET)")
        if not self.object_name:
            raise SheetFormatError("cell B1 must name the main object (e.g. ASSET)")
        if not self.action:
            raise SheetFormatError("cell C1 must name the action (Query, AddChange, Add, Change, Delete)")
        self.action = normalize_action(self.action)
        if not any(c.is_data for c in self.columns):
            raise SheetFormatError("row 2 has no attribute columns")
        seen: dict[str, int] = {}
        for column in self.columns:
            if column.header in seen:
                raise SheetFormatError(
                    f"duplicate header {column.header!r} in columns {seen[column.header] + 1} and {column.index + 1}"
                )
            seen[column.header] = column.index
        if sum(c.kind == "result" for c in self.columns) > 1:
            raise SheetFormatError("more than one [RESULT] column")

    # -- accessors -------------------------------------------------------------------
    @property
    def attr_columns(self) -> list[Column]:
        return [c for c in self.columns if c.is_data]

    @property
    def id_columns(self) -> list[Column]:
        return [c for c in self.columns if c.kind == "id"]

    @property
    def result_column(self) -> Column | None:
        return next((c for c in self.columns if c.kind == "result"), None)

    @property
    def error_column(self) -> Column | None:
        return next((c for c in self.columns if c.kind == "error"), None)

    @property
    def relation_paths(self) -> list[tuple[str, ...]]:
        """Every child relation chain named in the header, parents before
        children, in first-appearance order."""
        ordered: list[tuple[str, ...]] = []
        for column in self.attr_columns:
            for depth in range(1, len(column.path) + 1):
                prefix = column.path[:depth]
                if prefix not in ordered:
                    ordered.append(prefix)
        return ordered

    def row_number(self, data_index: int) -> int:
        return START_ROW + data_index

    def cell(self, data_index: int, column: Column) -> Any:
        row = self.data[data_index]
        return row[column.index] if column.index < len(row) else None

    def data_row_count(self) -> int:
        """Rows up to (not including) the first fully empty one — MxLoader stops
        at the first empty row too."""
        for i, row in enumerate(self.data):
            if all(v is None or str(v).strip() == "" for v in row):
                return i
        return len(self.data)

    # -- mutation (SetResult / SetId / AddResultCol / CleanSheet) -------------------------
    def ensure_result_column(self) -> Column:
        existing = self.result_column
        if existing:
            return existing
        column = Column(len(self.columns), RESULT_HEADER, "result")
        self.columns.append(column)
        return column

    def _set(self, row_number: int, column: Column, value: Any) -> None:
        row = self.data[row_number - START_ROW]
        while len(row) <= column.index:
            row.append(None)
        row[column.index] = value

    def set_result(self, row_number: int, text: str) -> None:
        column = self.result_column
        if column is not None:
            self._set(row_number, column, text)

    def set_error(self, row_number: int, text: str) -> None:
        """MxLoader's `[ERROR]` virtual attribute: only present if the user added
        it; holds the error for failed rows and is blank for successful ones, so
        a filter on it isolates the rows to fix and re-run."""
        column = self.error_column
        if column is not None:
            self._set(row_number, column, text)

    def set_id(self, row_number: int, attr: str, value: Any) -> None:
        attr = attr.lower()
        for column in self.id_columns:
            if not column.path and column.attr == attr:
                self._set(row_number, column, value)

    def clear_data(self) -> None:
        self.data = []

    # -- I/O -----------------------------------------------------------------------------
    def config_row(self) -> list[str]:
        return [self.object_service, self.object_name, self.action, self.where, self.order_by, self.oslc_directives]

    def header_row(self) -> list[str]:
        return [c.header for c in self.columns]

    def save(self, path: Path | None = None) -> Path:
        target = Path(path or self.path)
        if target.suffix.lower() in (".xlsx", ".xlsm"):
            _save_xlsx(self, target)
        else:
            _save_csv(self, target)
        return target


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value)


def _build(config: list[Any], header: list[Any], data: list[list[Any]], path: Path, sheet_name: str | None) -> MxSheet:
    config = [_cell_text(v).strip() for v in (list(config) + [""] * 6)[:6]]
    while header and (header[-1] is None or str(header[-1]).strip() == ""):
        header = header[:-1]
    if not header:
        raise SheetFormatError(f"{path}: row 2 (attribute header) is empty")
    columns = [parse_header(i, h) for i, h in enumerate(header)]
    return MxSheet(
        object_service=config[0],
        object_name=config[1],
        action=config[2],
        where=config[3],
        order_by=config[4],
        oslc_directives=config[5],
        columns=columns,
        data=[list(row) for row in data],
        path=path,
        sheet_name=sheet_name,
    )


def load_sheet(path: Path, sheet_name: str | None = None) -> MxSheet:
    """Load an MxLoader sheet from a .csv file, or one worksheet of an
    .xlsx/.xlsm workbook (the named one, else the active one)."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return _load_xlsx(path, sheet_name)
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    if len(rows) < HEADER_ROW:
        raise SheetFormatError(f"{path}: expected row 1 (config) and row 2 (attribute header)")
    return _build(rows[0], rows[1], rows[2:], path, None)


def _openpyxl():
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - exercised only without the extra
        raise SheetFormatError("Excel sheets need openpyxl: pip install 'maxmig[excel]'") from exc
    return openpyxl


def _load_xlsx(path: Path, sheet_name: str | None) -> MxSheet:
    workbook = _openpyxl().load_workbook(path, data_only=True)
    try:
        ws = workbook[sheet_name] if sheet_name else workbook.active
    except KeyError:
        raise SheetFormatError(f"{path}: no worksheet named {sheet_name!r} (have {', '.join(workbook.sheetnames)})") from None
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if len(rows) < HEADER_ROW:
        raise SheetFormatError(f"{path}[{ws.title}]: expected row 1 (config) and row 2 (attribute header)")
    return _build(rows[0], rows[1], rows[2:], path, ws.title)


def _save_csv(sheet: MxSheet, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".partial")
    width = len(sheet.columns)
    with tmp.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(sheet.config_row())
        writer.writerow(sheet.header_row())
        for row in sheet.data:
            writer.writerow([_cell_text(v) for v in (list(row) + [None] * width)[:width]])
    tmp.replace(path)


def _save_xlsx(sheet: MxSheet, path: Path) -> None:
    openpyxl = _openpyxl()
    if path.exists():
        workbook = openpyxl.load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
        title = sheet.sheet_name or workbook.active.title
        ws = workbook[title] if title in workbook.sheetnames else workbook.create_sheet(title)
    else:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.title = (sheet.sheet_name or sheet.object_name or "Sheet")[:31]
    if ws.max_row >= START_ROW:
        ws.delete_rows(START_ROW, ws.max_row - START_ROW + 1)
    for col, value in enumerate(sheet.config_row(), start=1):
        ws.cell(row=1, column=col, value=value or None)
    for col, value in enumerate(sheet.header_row(), start=1):
        ws.cell(row=HEADER_ROW, column=col, value=value)
    for offset, row in enumerate(sheet.data):
        for col, value in enumerate(row, start=1):
            if value is not None and value != "":
                ws.cell(row=START_ROW + offset, column=col, value=value)
    workbook.save(path)
